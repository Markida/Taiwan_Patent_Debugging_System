"""Create the customer-facing product introduction for the clean edition."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "deliverables"
OUTPUT_PATH = OUTPUT_DIR / "Saint-Island_Patent_MDS_無彩蛋版產品與檢核能力介紹.docx"
ICON_PATH = PROJECT_ROOT / "app" / "resources" / "app_icon.png"


NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
GRAY = "5B6573"
LIGHT_BLUE = "EAF3FA"
PALE_BLUE = "F3F8FC"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
RED = "B42318"
GOLD = "9A6700"
GREEN = "18794E"


def set_run_font(run, *, size=None, color=INK, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, (cell, width) in enumerate(zip(row.cells, widths_dxa)):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_paragraph_border(paragraph, *, color=BLUE, size=18, space=6, side="left"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)
    borders.append(border)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])
    tail = paragraph.add_run(" 頁")
    set_run_font(tail, size=9, color=GRAY)


def add_real_numbering(doc, style_name, *, bullet=False):
    styles = doc.styles
    try:
        style = styles[style_name]
    except KeyError:
        style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles["Normal"]
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.167

    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(level_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    first_num_index = next(
        (
            index
            for index, child in enumerate(numbering)
            if child.tag == qn("w:num")
        ),
        len(numbering),
    )
    numbering.insert(first_num_index, abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)

    p_style = style._element.get_or_add_pPr()
    num_pr = p_style.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_style.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_ref = OxmlElement("w:numId")
    num_ref.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_ref])
    return num_id


def apply_paragraph_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_ref = OxmlElement("w:numId")
    num_ref.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_ref])
    p_pr.append(num_pr)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    doc._product_number_num_id = add_real_numbering(
        doc, "Product Number", bullet=False
    )
    bullet_style = doc.styles["List Bullet"]
    bullet_style.font.name = "Calibri"
    bullet_style._element.get_or_add_rPr().rFonts.set(
        qn("w:eastAsia"), "Microsoft JhengHei"
    )
    bullet_style.font.size = Pt(11)
    bullet_style.paragraph_format.space_after = Pt(8)
    bullet_style.paragraph_format.line_spacing = 1.167

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    left = header_p.add_run("Saint-Island_Patent_MDS")
    set_run_font(left, size=8.5, color=GRAY, bold=True)
    right = header_p.add_run("    產品介紹｜無彩蛋乾淨版")
    set_run_font(right, size=8.5, color=GRAY)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    add_page_number(footer_p)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_body(doc, text, *, bold_lead="", color=INK, italic=False):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, size=11, color=color, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, size=11, color=color, italic=italic)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=color, italic=italic)
    return p


def add_bullet(doc, text, *, bold_lead=""):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead and text.startswith(bold_lead):
        run = p.add_run(bold_lead)
        set_run_font(run, size=11, bold=True)
        run = p.add_run(text[len(bold_lead):])
        set_run_font(run, size=11)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11)
    return p


def add_number(doc, text, *, bold_lead=""):
    p = doc.add_paragraph(style="Product Number")
    apply_paragraph_numbering(p, doc._product_number_num_id)
    if bold_lead and text.startswith(bold_lead):
        run = p.add_run(bold_lead)
        set_run_font(run, size=11, bold=True)
        run = p.add_run(text[len(bold_lead):])
        set_run_font(run, size=11)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11)
    return p


def add_callout(doc, label, text, *, color=BLUE, fill=PALE_BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(10)
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    add_paragraph_border(p, color=color, size=24, space=7)
    lead = p.add_run(f"{label}　")
    set_run_font(lead, size=11, color=color, bold=True)
    body = p.add_run(text)
    set_run_font(body, size=11, color=INK)
    return p


def style_table(table, header=True):
    table.style = "Table Grid"
    if header:
        set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            if row_index == 0 and header:
                set_cell_shading(cell, LIGHT_GRAY)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    set_run_font(run, size=10.2, bold=(row_index == 0 and header))


def add_capability_table(doc):
    rows = [
        ("功能層", "主要工作", "客戶能直接看見的價值"),
        ("文件偵錯", "讀取完整 DOCX，檢查章節、段號、標點、符號、請求項與文字一致性。", "在送件或內部複核前，集中呈現高風險文字與結構問題。"),
        ("圖式標號識別", "從圖片或 PDF 辨識數字、英文及 prime 標號，並與文件符號清單比對。", "降低人工逐張抄錄及核對圖式標號的時間。"),
        ("段落圖式比對", "理解實施方式參閱的圖號，將段落元件標號與對應圖式 OCR 結果交叉比對。", "優先指出「文件寫到了，但指定圖式沒有」的缺漏。"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    set_table_geometry(table, [1800, 3360, 4200])
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    style_table(table)
    return table


def add_severity_table(doc):
    rows = [
        ("等級", "代表意義", "典型情況"),
        ("錯誤", "高度可能影響文件完整性或一致性，應優先確認。", "缺少章節、段號不連續、請求項依附無效、元件標號錯配、圖式缺失標號。"),
        ("警告", "可能是問題，也可能是合法個案，需要專業人員判讀。", "疑似元件名稱錯字、標的名稱不一致、表格中的跨欄請求項。"),
        ("資訊", "格式正規化或提醒，不直接代表法律內容有誤。", "全形英數字、非標準 prime mark、特定用語人工確認。"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    set_table_geometry(table, [1250, 3000, 5110])
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    style_table(table)
    for row, fill in zip(table.rows[1:], ("FDECEC", "FFF4D6", "EAF3FA")):
        set_cell_shading(row.cells[0], fill)
    return table


def add_cover(doc):
    for _ in range(2):
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(8)
    if ICON_PATH.is_file():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run()
        shape = run.add_picture(str(ICON_PATH), width=Inches(0.88))
        doc_pr = shape._inline.docPr
        doc_pr.set("descr", "Saint-Island_Patent_MDS 應用程式圖示")
        p.paragraph_format.space_after = Pt(18)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(5)
    run = kicker.add_run("PATENT MISTAKE DETECTION SYSTEM")
    set_run_font(run, size=10.5, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("Saint-Island_Patent_MDS")
    set_run_font(run, size=28, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(22)
    run = subtitle.add_run("專利文件與圖式一致性檢核平台")
    set_run_font(run, size=15, color=DARK_BLUE, bold=True)

    desc = doc.add_paragraph()
    desc.paragraph_format.space_after = Pt(20)
    run = desc.add_run(
        "將專利說明書的文字規則、圖式標號 OCR 與段落圖式關聯整合於同一套工具，協助專業人員在送件前更快發現遺漏、矛盾與誤植。"
    )
    set_run_font(run, size=12, color=INK)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(12)
    meta.paragraph_format.space_after = Pt(4)
    run = meta.add_run("版本：v2.1.07c 公司標準版")
    set_run_font(run, size=10.5, color=GRAY, bold=True)
    meta2 = doc.add_paragraph()
    meta2.paragraph_format.space_after = Pt(24)
    run = meta2.add_run(f"文件性質：產品與檢核能力介紹（非操作手冊）｜{date.today():%Y-%m-%d}")
    set_run_font(run, size=10, color=GRAY)

    add_callout(
        doc,
        "核心定位",
        "系統不替代專利專業判斷，也不直接改寫原始 Word；它的角色是先把最值得人工注意的位置找出來，讓複核工作更集中、更可追溯。",
        color=BLUE,
        fill=LIGHT_BLUE,
    )
    doc.add_page_break()


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "Saint-Island_Patent_MDS 無彩蛋版產品與檢核能力介紹"
    props.subject = "專利文件偵錯、圖式標號 OCR 與段落圖式比對產品介紹"
    props.author = "Saint-Island Patent MDS"
    props.keywords = "專利文件, 錯誤檢核, OCR, 圖式標號, 請求項"

    add_cover(doc)

    add_heading(doc, "一、產品概要", 1)
    add_body(
        doc,
        "Saint-Island_Patent_MDS 是一套面向專利文件品質控管的桌面應用程式。無彩蛋乾淨版保留所有正式工作功能，將完整專利說明書、圖式 PDF／圖片與符號清單放在同一個檢核流程中，讓使用者不必分別以人工方式逐頁、逐段與逐項交叉核對。",
    )
    add_body(
        doc,
        "系統採用「確定性文字規則＋圖像辨識＋跨資料來源比對」的設計：可明確判斷的格式與邏輯問題列為錯誤；可能涉及個案寫法的情況列為警告；正規化建議則列為資訊。所有結果均回到原始段落定位，由專業人員決定是否修訂。",
    )
    add_callout(
        doc,
        "客戶可理解成",
        "它是一位先做全面初篩的文件複核助手：先找出章節、請求項、元件、標號與圖式之間的不一致，再把人的時間留給真正需要判斷的內容。",
        color=GREEN,
        fill="EDF8F2",
    )

    add_heading(doc, "二、三大核心功能", 1)
    add_capability_table(doc)
    add_body(
        doc,
        "三個功能共用同一份案件資料。文件偵錯頁可擷取完整符號說明與代表圖符號說明；圖式標號頁可直接選擇其中一份清單進行比對；完成 OCR 後，段落圖式比對頁可沿用文件與圖片結果，形成由文字到圖式的連續檢核鏈。",
    )

    add_heading(doc, "三、專利文件偵錯：從結構到語意一致性", 1)
    add_body(
        doc,
        "文件偵錯功能針對正式 DOCX 進行唯讀解析，不會直接覆寫或輸出修訂後的 Word。系統保留原始段落與字元位置，將問題依錯誤、警告及資訊排序，並在選取結果時顯示完整原文段落及醒目的錯誤詞語。",
    )

    add_heading(doc, "3.1 文件類型、章節與基本版面", 2)
    add_bullet(doc, "辨認發明或新型文件；一旦類型確定，只套用該類型的章節與用語規則。")
    add_bullet(doc, "確認三大章節——摘要、說明書、申請專利範圍——各出現一次，且順序正確。")
    add_bullet(doc, "確認中型章節位於正確的大章節，並依正式順序排列；偵測重複、缺漏、空內容或不在固定清單內的中括號標題。")
    add_bullet(doc, "說明書及申請專利範圍前應使用真正的 Word『下一頁分節符號』，而非以空白段落模擬換頁。")
    add_bullet(doc, "章節標題只比較文字內容，不以字型、字體大小、括號或空白差異判定錯誤。")
    add_bullet(doc, "檢查 A4 頁面、邊界、頁首與頁尾距離是否符合正式範本；正文文字大小不列入檢核。")
    add_bullet(doc, "摘要與說明書中的中文名稱必須逐字一致；英文名稱可省略，但若使用，兩處必須同時存在且內容一致。")

    add_heading(doc, "3.2 說明書段號與段落完整性", 2)
    add_bullet(doc, "檢查說明書第二大章節中的正文段號是否由 1 開始並依序連續。")
    add_bullet(doc, "接受【1】、【01】、【001】及【0001】等一至四位數顯示；補零方式不影響合法性。")
    add_bullet(doc, "辨認手動鍵入的段號與 Word 自動編號，並將手動段號、空段號、段號跳號及段號出現在錯誤章節列為錯誤。")
    add_bullet(doc, "符號說明有自己的清單語意，不套用一般說明書段號格式；圖式簡單說明的圖說亦依專用規則處理。")
    add_callout(
        doc,
        "目前刻意不檢查",
        "段落中的異常空白不列為錯誤，以避免不同 Word 排版習慣造成大量無效結果。",
        color=GOLD,
        fill="FFF8E7",
    )

    add_heading(doc, "3.3 段落結尾與標點邏輯", 2)
    add_bullet(doc, "中文摘要段落應以全形句號結尾。")
    add_bullet(doc, "說明書的數字段落原則上以全形句號結尾；實施方式的非末段可用全形冒號引出後續內容。")
    add_bullet(doc, "實施方式若為表格內容，不因表格段落缺少結尾標點而報錯。")
    add_bullet(doc, "圖式簡單說明的導言應以全形冒號結尾；中間圖說使用全形分號、倒數第二筆使用『；及』、最後一筆使用全形句號。")
    add_bullet(doc, "申請專利範圍中的請求項以句號結束；跨行獨立項另檢查首行冒號、中間分號、倒數第二行『；及』及末行句號。")
    add_bullet(doc, "符號說明不檢查結尾標點。")

    add_heading(doc, "四、符號說明、元件標號與文字錯誤", 1)
    add_heading(doc, "4.1 完整符號說明與代表圖符號說明", 2)
    add_bullet(doc, "擷取數字、英文字母及其他可辨識字元的元件標號；標號不限定為純數字。")
    add_bullet(doc, "可展開 S1～S8 等連續範圍，供後續 OCR 清單比對使用。")
    add_bullet(doc, "檢查符號是否缺少元件名稱、是否重複列出，或同一符號是否對應不同名稱。")
    add_bullet(doc, "確認代表圖所列符號已存在於完整符號說明，並比較兩處的元件名稱是否一致。")
    add_bullet(doc, "不強制檢查符號說明的數字排版、冒號形式或分隔格式，重點放在符號與名稱內容。")

    add_heading(doc, "4.2 實施方式中的元件名稱與標號", 2)
    add_body(
        doc,
        "元件標號檢查只作用於『實施方式』。系統優先選擇最長且完整的元件名稱，例如同時存在『齒輪』與『第一齒輪』時，以『第一齒輪』為完整名稱，避免把 11 誤拆成齒輪的標號 1。",
    )
    add_bullet(doc, "元件後方標號若未列入完整符號說明，提出提醒；已登錄元件若使用錯誤標號，列為錯誤。")
    add_bullet(doc, "只有前方具『該、該等、每一該、其中一該』或數量詞等明確元件指稱時，才檢查後方是否缺少標號，降低一般敘述中的誤報。")
    add_bullet(doc, "支援『發光二極體（LED）512』等括號補充名稱後再接標號的寫法。")
    add_bullet(doc, "案件標的名稱、實施例名稱及已證明為實施方式限定的局部元件，會依情境排除不適用的標號錯誤。")
    add_bullet(doc, "使用者可為單一文件建立『實施方式段落出現之無標號元件』清單；完整詞語中的短元件名稱也會一併受到保護。")

    add_heading(doc, "4.3 疑似錯字與自訂文字規則", 2)
    add_bullet(doc, "檢查與符號說明元件名稱僅差一個中文字的疑似誤植，並排除主／副、完整合法重疊名稱及白名單詞語等已知合法差異。")
    add_bullet(doc, "相似詞比對也會考慮元件名稱字序重排後只差一字的情況，用於發現不易察覺的輸入錯誤。")
    add_bullet(doc, "任意兩個以上連續且相同的中文字（例如『的的』）列為錯誤。")
    add_bullet(doc, "共用黑名單可新增特定錯誤文字；共用白名單可排除不應參與相似詞警告的合法詞語，重新開啟程式後仍可沿用。")
    add_bullet(doc, "全形英數字、非標準 prime mark 及特定用語以資訊等級提示，不直接宣告內容錯誤。")
    add_bullet(doc, "發明文件出現『本新型』或新型文件出現『本發明』等混用情況，列為錯誤。")

    add_heading(doc, "五、申請專利範圍的重點檢核", 1)
    add_body(
        doc,
        "請求項是本系統規則最密集的區域。程式會先依『請求項 N』界定每一項的範圍，優先擷取最初出現的項次邊界，因此可處理獨立項換行、後文再次引用其他請求項，以及一份案件具有多個獨立標的的情況。",
    )

    add_heading(doc, "5.1 項次、依附關係與句型", 2)
    add_bullet(doc, "檢查請求項是否有明確項次、是否使用 Word 自動編號、是否由 1 起連續、以及是否有空白請求項。")
    add_bullet(doc, "附屬項只能引用已存在且編號較小的請求項；支援單項、範圍及多項依附，例如『請求項1至8中任一項』。")
    add_bullet(doc, "多項附屬項應具有『任一項』或『或』等選擇式文字，且不得直接或間接依附另一多項附屬項。")
    add_bullet(doc, "檢查一個請求項原則上是否為一個完整句子，以及跨行獨立項的各行標點。")
    add_bullet(doc, "比較附屬項與所依附請求項的標的名稱；若系統無法可靠判別，改列警告而非自行推定。")
    add_bullet(doc, "後續以『一種……的控制方法』或與請求項1相同的分行格式開始時，可辨認為第二獨立項，並保留多個標的名稱。")
    add_bullet(doc, "『二如請求項1所述的……』『複數如請求項1所述的……』等量化引用，也可作為新獨立項及依附來源判斷的一部分。")

    add_heading(doc, "5.2 元件第一次揭露與後續指稱", 2)
    add_body(
        doc,
        "底層邏輯以『冠詞／數量詞＋可能的修飾子句＋的／之＋元件名稱』為主要解析骨架。這使系統不只看元件前一個字，而能處理『一連通該底部開口的電池空間』或『一沿一頂底方向設置的基座』等長修飾語句。",
    )
    add_bullet(doc, "元件第一次出現時應有數量詞或其他建立先行基礎的表達；未建立即使用『該／該等』列為錯誤。")
    add_bullet(doc, "單數可由『一、至少一』等建立；複數可由『二／兩、三、數個、多個、複數』等建立，並支援指定數量。")
    add_bullet(doc, "後續指稱會核對單複數：單數通常使用『該』，複數可使用『該等、每一該、其中一該、任意該、其中複數該、各自的該』等合法形式。")
    add_bullet(doc, "『還包括另一該開槽』等新增另一成員的語句，會使該元件在後文具備複數基礎。")
    add_bullet(doc, "『每一該架體形成有複數橫梁』等分配式關係，會把每一母體所具有的子元件正確視為整體複數。")
    add_bullet(doc, "在『每一該送料輥具有一輥本體……』的限定範圍內，子元件可暫時以單數指稱；離開該限定關係後仍依整體單複數基礎判斷。")
    add_bullet(doc, "多位階元件如『該安裝座的底壁』或『每一該送料孔道的該第一孔段』，後位階元件不被強迫重複加入新的冠詞。")
    add_bullet(doc, "『步驟』屬特殊元件，只嚴格檢查其標號，不對單複數冠詞做過度判斷。")
    add_callout(
        doc,
        "為何這部分重要",
        "請求項中最難人工穩定發現的問題，往往不是單一錯字，而是某個元件在前文尚未建立、前後單複數改變，或跨請求項依附後失去先行基礎。系統會把這些關係轉成可逐項查看的錯誤。",
        color=RED,
        fill="FDECEC",
    )

    add_heading(doc, "5.3 請求項中的表格", 2)
    add_body(
        doc,
        "表格儲存格的文字仍會逐格接受重複字、自訂黑名單及其他文字規則檢查；但系統不會跨欄拼接請求項語法，因此若申請專利範圍以表格呈現，會提出警告並要求人工確認結構。這項限制可避免程式因表格讀取順序而產生不可靠結論。",
    )

    add_heading(doc, "六、圖式標號 OCR 與清單比對", 1)
    add_body(
        doc,
        "圖式標號識別功能處理多張圖片或完整 PDF。PDF 會以 300 DPI 轉為圖片，再由正式 ONNX／OCR 模型偵測標號位置與內容。辨識支援 0–9、英文字母大小寫及 prime mark，可將相鄰字元組合為 10A、7' 等完整標號。",
    )
    add_bullet(doc, "批次辨識多頁 PDF 或多張圖片，結果以 Pic_01、Pic_02 等固定名稱整理。")
    add_bullet(doc, "預設信心檢視門檻為 0.60；低於門檻的標號會以紅框與優先排序呈現，方便人工集中確認。")
    add_bullet(doc, "使用者可直接修改、增加或刪除辨識標號；點選圖片辨識框可對應選取清單項目。")
    add_bullet(doc, "支援滑鼠滾輪與按鈕縮放、適合視窗、單張左右旋轉及全部圖片右旋，以便閱讀細小標號。")
    add_bullet(doc, "每一圖片頁可設定一個或多個圖號，例如圖1與圖2同頁，亦支援圖3A、圖3B。")
    add_bullet(doc, "可在『完整符號說明』與『代表圖符號說明』之間切換，選擇實際用來比對的標號清單。")
    add_bullet(doc, "比對結果同時提供 All Pictures 全體缺漏，以及每張圖片個別的『清單有、圖片沒有』與『圖片有、清單沒有』。")
    add_bullet(doc, "標號可依逐字元邏輯排序，讓 1、121、131、2、21、3 等結果更符合人工查找習慣。")
    add_callout(
        doc,
        "設計原則",
        "OCR 採取『先提高召回、再讓低信心結果醒目可修正』的方式。系統不會把模型輸出視為不可更動的答案，而是把人工確認納入正式流程。",
        color=BLUE,
        fill=LIGHT_BLUE,
    )

    add_heading(doc, "七、實施方式與圖式比對", 1)
    add_body(
        doc,
        "第三個功能把文件中的實施方式段落與 OCR 圖片並排呈現。系統解析段落正在參閱哪些圖式，擷取該段中『元件名稱＋符號清單標號』的組合，再確認這些標號是否出現在對應圖式。其主要結果名稱為『圖式缺失標號』，並同時顯示標號所代表的中文元件名稱。",
    )

    add_heading(doc, "7.1 圖號語句的理解能力", 2)
    add_bullet(doc, "支援參閱、參照、參考、參見、詳見、見、如、在、於、由等常見提示詞。")
    add_bullet(doc, "支援單圖、多圖、連續範圍與混合寫法，例如『參閱圖1』『參照圖1、圖2及圖3』『參考圖1至4』『參閱圖3A、3B及4』。")
    add_bullet(doc, "圖號不必位於段落開頭；『在圖3中……』『圖1和圖2分別示出……以及圖3……』等段中切換也能建立新的圖式範圍。")
    add_bullet(doc, "沒有新圖號的段落可沿用前一段參閱圖式；出現『例如為圖5』等補充案例時，後文同時保留原圖與新增圖的依附關係。")
    add_bullet(doc, "同一 Word 數字段落中的多個實體段落會合併為一個邏輯段落，避免同一內容被拆成多列。")
    add_bullet(doc, "當實施方式出現『綜上所述』時，包含該段及其後內容不再進行段落圖式比對。")

    add_heading(doc, "7.2 標號比對如何降低誤報", 2)
    add_bullet(doc, "只比對完整符號說明中確實存在的標號，且標號前必須辨認到相對應元件名稱；單獨出現的數字不當成元件標號。")
    add_bullet(doc, "比對採單向原則：段落提到而圖式缺少的標號會報錯；圖式多出但段落沒有提到的標號暫不報錯。")
    add_bullet(doc, "若一張圖片同時包含多個圖號，採頁面子集合比對，避免把圖1與圖2同頁誤判成兩張獨立圖片。")
    add_bullet(doc, "元件後的『（見圖1）』『（如圖3所示）』『（參閱圖3）』『（參考圖3）』『（參圖3）』可將該元件局部綁定到補充圖式，不會錯誤切換整段的主要參閱圖。")
    add_bullet(doc, "同一小段落中已建立的補充圖式綁定可沿用到該元件後續出現的位置。")
    add_bullet(doc, "若文件參閱的圖號尚未載入 OCR，系統會明確標示缺少哪一張圖，而非把所有標號一律判成圖式缺失。")

    add_heading(doc, "7.3 剖視圖與整體圖式使用檢查", 2)
    add_bullet(doc, "當圖式簡單說明出現剖視／剖面／截面／斷面及 VI–VI 等羅馬剖切線時，檢查左右羅馬數字是否一致且為合法羅馬數字。")
    add_bullet(doc, "將羅馬數字換算為阿拉伯數字，確認例如 VI 對應圖6；並檢查來源圖 OCR 是否真的辨識到 VI。")
    add_bullet(doc, "若圖說無法辨認來源圖、來源圖尚未載入，或來源圖沒有該剖切線，分別提供具體錯誤原因。")
    add_bullet(doc, "檢查每一張已載入 OCR 的圖式是否至少在全文引用、參閱或圖式簡單說明中出現一次。")

    section_eight = add_heading(doc, "八、結果呈現與可追溯性", 1)
    section_eight.paragraph_format.page_break_before = True
    add_severity_table(doc)
    add_body(
        doc,
        "文件問題會優先以數字段號定位；若沒有段號，則使用中型章節定位。選取錯誤或警告時，系統顯示完整原始 Word 段落，並只將真正出錯的詞語標成紅色，避免使用者只看到脫離語境的短句。請求項問題則顯示整個請求項內容，方便判斷跨行或先行揭露關係。",
    )
    add_body(
        doc,
        "所有修改仍由使用者回到原始 Word 完成。這種唯讀設計讓檢核結果與正式文件分離，可避免自動改寫破壞段落編號、分節符號、欄位或其他申請文件格式。",
    )

    add_heading(doc, "九、適合的導入情境", 1)
    add_number(doc, "送件前總檢：在正式提交前，快速確認章節、標點、請求項、元件及圖式之間是否有明顯遺漏。")
    add_number(doc, "撰稿後交叉複核：讓不同承辦人以一致規則查看同一份文件，降低只依個人記憶檢查的差異。")
    add_number(doc, "大量圖式案件：利用 OCR 與完整／代表圖符號清單，縮短逐頁尋找標號的時間。")
    add_number(doc, "修改稿回查：文件或圖式經過多次修訂後，再次確認名稱、標號與參閱圖號沒有因局部改動而失去一致性。")
    add_number(doc, "內部品質標準落地：透過共用黑名單、白名單及單文件白名單，把常見錯誤與案件例外累積成可重複使用的檢核知識。")

    section_ten = add_heading(doc, "十、產品界線與專業覆核", 1)
    section_ten.paragraph_format.page_break_before = True
    add_body(
        doc,
        "Saint-Island_Patent_MDS 是錯誤發現與一致性檢核工具，不是法律意見產生器，也不判斷發明是否具備新穎性、進步性或專利要件。以下情況仍應由專業人員覆核：",
    )
    add_bullet(doc, "OCR 低信心、線條干擾、極小字元或特殊字型造成的辨識不確定性。")
    add_bullet(doc, "合法但罕見的請求項語法、複雜表格、跨欄內容或案件特有用語。")
    add_bullet(doc, "同一名詞在不同技術語境中的語意差異，以及需要專利法專業判斷的實質內容。")
    add_bullet(doc, "段落圖式比對已涵蓋多種參閱語法與例外，結果仍建議視為高效率初篩。")
    add_callout(
        doc,
        "無彩蛋乾淨版範圍",
        "本版本僅包含文件偵錯、圖式標號識別、段落圖式比對及其必要資料交換功能；不含聊天室、貪食蛇、雙人彈球或其他隱藏功能。",
        color=NAVY,
        fill=LIGHT_GRAY,
    )

    add_heading(doc, "結語", 1)
    add_body(
        doc,
        "專利文件的風險通常不是單一錯字，而是章節、請求項、元件名稱、符號與圖式之間長距離的不一致。Saint-Island_Patent_MDS 將這些原本分散的檢查集中為可定位、可說明、可人工確認的結果，協助專業人員把時間投入真正需要判斷的地方。",
    )
    final = doc.add_paragraph()
    final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    final.paragraph_format.space_before = Pt(18)
    run = final.add_run("Saint-Island_Patent_MDS｜Mistake Detection System")
    set_run_font(run, size=10.5, color=BLUE, bold=True)

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_document())
